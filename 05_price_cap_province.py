"""
Quota provinciale degli impianti delle bandiere che applicano il price cap (Agip Eni, Api-Ip, Q8)
sul totale degli impianti stradali (esclusi gli autostradali).

Fonte: anagrafica_impianti_attivi.csv (estrazione del 28/09/2026), nella cartella principale.

Province: si usa l'assetto Istat attuale (110 province/UTS, file in riferimenti/).
Il MIMIT usa ancora le vecchie sigle sarde (SS, NU, OR, CA, SU), che non corrispondono più
ai confini attuali: per la Sardegna ogni impianto è assegnato alla provincia tramite il codice
Istat del suo comune. Per il resto d'Italia si usa la sigla MIMIT, abbinata al codice Istat
della provincia a cui appartiene la maggioranza dei comuni degli impianti con quella sigla.

Aggiunge (o sostituisce) il foglio "Price cap per provincia" nel file Excel prodotto da
04_costruisci_excel.py e salva anche output/price_cap_province.csv.
"""
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook

from comune import BASE, DIR_OUTPUT, FILE_XLSX, leggi_anagrafica, scrivi_foglio

FILE_ANAGRAFICA = BASE / "anagrafica_impianti_attivi.csv"
DIR_RIFERIMENTI = BASE / "riferimenti"
BANDIERE_PRICE_CAP = ["Agip Eni", "Api-Ip", "Q8"]
SIGLE_SARDEGNA = {"SS", "NU", "OR", "CA", "SU"}
UTS_SARDEGNA = {"312", "113", "114", "115", "116", "117", "318", "119"}
NOME_FOGLIO = "Price cap per provincia"


def normalizza(nome):
    """'Tortolì' e "TORTOLI'" -> 'TORTOLI'"""
    nome = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]", "", nome)


def carica_riferimenti():
    province = pd.read_csv(DIR_RIFERIMENTI / "province_istat.csv", dtype=str)
    comuni = pd.read_csv(DIR_RIFERIMENTI / "comuni_istat.csv", dtype=str)
    codici_uts = set(province["codice_istat"])

    # Il codice comune inizia con il codice della provincia "storica"; per le città metropolitane
    # il codice UTS è 200 + quello storico (es. Torino 001 -> 201, Cagliari 118 -> 318).
    def uts(codice_comune):
        storico = int(codice_comune[:3])
        return f"{storico + 200:03d}" if f"{storico + 200:03d}" in codici_uts else f"{storico:03d}"

    comuni["uts"] = comuni["codice_istat"].map(uts)
    comuni["chiave"] = comuni["denominazione"].map(normalizza)
    assert set(comuni["uts"]) <= codici_uts
    return province, comuni


def assegna_provincia(impianti, comuni):
    """Aggiunge a ogni impianto il codice Istat (UTS) della provincia attuale."""
    abb = impianti.merge(comuni[["chiave", "uts"]], left_on=impianti["Comune"].map(normalizza),
                         right_on="chiave", how="left")

    # Resto d'Italia: sigla MIMIT -> provincia prevalente fra i comuni abbinati
    fuori = abb[~abb["Provincia"].isin(SIGLE_SARDEGNA)].dropna(subset=["uts"])
    conta = fuori.groupby(["Provincia", "uts"]).size()
    prevalente = conta.groupby(level=0).idxmax().map(lambda t: t[1])
    quota = (conta.groupby(level=0).max() / conta.groupby(level=0).sum()).round(3)
    print("Sigle con abbinamento comune-provincia sotto il 95%:\n", quota[quota < 0.95].to_string() or "nessuna")
    assert prevalente.is_unique, "Due sigle MIMIT assegnate alla stessa provincia Istat"

    # Sardegna: provincia dal comune (in caso di omonimi, si prende il comune sardo)
    sard = abb[abb["Provincia"].isin(SIGLE_SARDEGNA) & abb["uts"].isin(UTS_SARDEGNA)]
    sard_uts = sard.drop_duplicates("idImpianto").set_index("idImpianto")["uts"]

    impianti = impianti.copy()
    impianti["uts"] = impianti["Provincia"].map(prevalente)
    sarde = impianti["Provincia"].isin(SIGLE_SARDEGNA)
    impianti.loc[sarde, "uts"] = impianti.loc[sarde, "idImpianto"].map(sard_uts)
    mancanti = impianti[impianti["uts"].isna()]
    assert mancanti.empty, f"Impianti senza provincia:\n{mancanti}"
    return impianti


def main():
    province, comuni = carica_riferimenti()
    ana = leggi_anagrafica(FILE_ANAGRAFICA)
    print(f"Impianti in anagrafica: {len(ana)}  |  autostradali esclusi: {(ana['Tipo Impianto'] == 'Autostradale').sum()}")
    ana = ana[ana["Tipo Impianto"] == "Stradale"].drop_duplicates("idImpianto")
    ana = assegna_provincia(ana, comuni)

    tab = pd.crosstab(ana["uts"], ana["Bandiera"])
    out = pd.DataFrame({"Impianti totali": tab.sum(axis=1)})
    for b in BANDIERE_PRICE_CAP:
        out[b] = tab[b] if b in tab else 0
    out["Impianti con price cap"] = out[BANDIERE_PRICE_CAP].sum(axis=1)
    out["Quota price cap (%)"] = (out["Impianti con price cap"] / out["Impianti totali"] * 100).round(1)

    nomi = province.set_index("codice_istat")["denominazione"]
    out = out.reset_index().rename(columns={"uts": "Codice Istat provincia"})
    out.insert(0, "Provincia", out["Codice Istat provincia"].map(nomi))
    out = out.sort_values(["Quota price cap (%)", "Provincia"], ascending=[False, True])
    assert len(out) == len(province), f"Province nel risultato: {len(out)} su {len(province)}"

    out.to_csv(DIR_OUTPUT / "price_cap_province.csv", index=False)
    wb = load_workbook(FILE_XLSX)
    if NOME_FOGLIO in wb.sheetnames:
        del wb[NOME_FOGLIO]
    scrivi_foglio(wb.create_sheet(NOME_FOGLIO), out,
                  ["testo", "codice", "intero", "intero", "intero", "intero", "intero", "percentuale"])
    wb.save(FILE_XLSX)

    tot = out[["Impianti totali", "Impianti con price cap"]].sum()
    print(f"\nItalia: {tot['Impianti con price cap']} impianti con price cap su {tot['Impianti totali']} "
          f"({tot['Impianti con price cap'] / tot['Impianti totali'] * 100:.1f}%)")
    print(out.head(10).to_string(index=False))
    print("...")
    print(out.tail(10).to_string(index=False))
    print(f"\nAggiunto il foglio '{NOME_FOGLIO}' a {FILE_XLSX}")


if __name__ == "__main__":
    main()
