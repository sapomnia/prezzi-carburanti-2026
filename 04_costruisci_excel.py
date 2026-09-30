"""
Costruisce il file Excel con i prezzi medi giornalieri di benzina e diesel (self service)
nel primo semestre 2026, impianti stradali (esclusi gli autostradali).

Regole applicate al dataset prodotto da 02_prepara_dati.py (vedi comune.carica_dati):
- solo prezzi self service (isSelf == 1);
- scartati i prezzi fuori dall'intervallo 1-3 euro/litro (errori di inserimento);
- media semplice tra gli impianti (ogni impianto pesa uno);
- operatori: 6 bandiere tenute distinte, grande distribuzione raggruppata, tutto il resto in "Altri operatori".

Fogli:
1. "Media nazionale": Data | Benzina (€/l) | Diesel (€/l)
2. "Per operatore":   Data | Operatore | Prodotto | Prezzo medio (€/l)   (formato verticale)
3. "Massimi per operatore": Operatore | Prodotto | Prezzo medio massimo (€/l) | Data
   cioè il picco della media giornaliera di ciascun operatore nel semestre, con il giorno in cui è stato
   raggiunto (in caso di parità, il primo giorno). Non usiamo il prezzo del singolo impianto perché i valori
   più alti sono quasi sempre prezzi di comodo (es. 2,999 €/l) e non prezzi realmente praticati.

Output: output/prezzi_medi_carburanti_S1_2026.xlsx (+ log dei prezzi scartati)
"""
import pandas as pd
from openpyxl import Workbook

from comune import DIR_OUTPUT, FILE_XLSX, ORDINE_OPERATORI, ORDINE_PRODOTTI, carica_dati, scrivi_foglio


def main():
    DIR_OUTPUT.mkdir(exist_ok=True)
    d = carica_dati()

    # Foglio 1: media nazionale giornaliera
    nazionale = (d.groupby(["data", "prodotto"], observed=True)["prezzo"].mean()
                 .unstack("prodotto")[ORDINE_PRODOTTI].round(3).reset_index())
    nazionale.columns = ["Data", "Benzina (€/l)", "Diesel (€/l)"]

    # Foglio 2: media giornaliera per operatore, formato verticale
    per_op = (d.groupby(["data", "operatore", "prodotto"], observed=True)["prezzo"].mean()
              .round(3).reset_index())
    per_op["operatore"] = pd.Categorical(per_op["operatore"], ORDINE_OPERATORI, ordered=True)
    per_op["prodotto"] = pd.Categorical(per_op["prodotto"].astype(str), ORDINE_PRODOTTI, ordered=True)
    per_op = per_op.sort_values(["data", "operatore", "prodotto"])
    per_op.columns = ["Data", "Operatore", "Prodotto", "Prezzo medio (€/l)"]

    # Foglio 3: picco della media giornaliera per operatore e prodotto (primo giorno in caso di parità)
    idx = per_op.groupby(["Operatore", "Prodotto"], observed=True)["Prezzo medio (€/l)"].idxmax()
    massimi = per_op.loc[idx, ["Operatore", "Prodotto", "Prezzo medio (€/l)", "Data"]]
    massimi.columns = ["Operatore", "Prodotto", "Prezzo medio massimo (€/l)", "Data"]
    massimi = massimi.sort_values(["Operatore", "Prodotto"])
    pari = per_op.merge(massimi, left_on=["Operatore", "Prodotto", "Prezzo medio (€/l)"],
                        right_on=["Operatore", "Prodotto", "Prezzo medio massimo (€/l)"])
    print("Giorni in cui è stato toccato il massimo:\n",
          pari.groupby(["Operatore", "Prodotto"], observed=True).size().to_string())

    # Controlli: 181 giorni, nessun buco
    giorni = pd.date_range("2026-01-01", "2026-06-30")
    assert list(nazionale["Data"]) == list(giorni), "Giorni mancanti nel foglio 1"
    assert len(per_op) == len(giorni) * len(ORDINE_OPERATORI) * len(ORDINE_PRODOTTI), \
        "Combinazioni data/operatore/prodotto mancanti nel foglio 2"

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Media nazionale"
    scrivi_foglio(ws1, nazionale, ["data", "prezzo", "prezzo"])
    ws2 = wb.create_sheet("Per operatore")
    scrivi_foglio(ws2, per_op, ["data", "testo", "testo", "prezzo"])
    ws3 = wb.create_sheet("Massimi per operatore")
    scrivi_foglio(ws3, massimi, ["testo", "testo", "prezzo", "data"])
    wb.save(FILE_XLSX)

    # Riepilogo: impianti self medi al giorno per operatore e prodotto
    impianti = (d.groupby(["operatore", "prodotto", "data"], observed=True)["idImpianto"].nunique()
                .groupby(level=[0, 1], observed=True).mean().round(0).unstack("prodotto"))
    print("\nImpianti self medi al giorno:\n", impianti.reindex(ORDINE_OPERATORI))
    print(f"\nScritto {FILE_XLSX}")


if __name__ == "__main__":
    main()
