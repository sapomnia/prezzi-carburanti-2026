"""
Funzioni condivise: individuazione dei file e lettura robusta dei CSV MIMIT.

Formato dei file (separatore "|"):
- riga 1: "Estrazione del AAAA-MM-GG"
- riga 2: intestazione
- anagrafica: idImpianto|Gestore|Bandiera|Tipo Impianto|Nome Impianto|Indirizzo|Comune|Provincia|Latitudine|Longitudine
- prezzi:     idImpianto|descCarburante|prezzo|isSelf|dtComu

Nell'anagrafica alcuni "Nome Impianto" e, più raramente, "Gestore" contengono a loro volta il
carattere "|" (es. "STOIL SIMPLE | gestori.prezzibenzina.it"): leggiamo quindi le righe a mano.
Gli ultimi 4 campi si prendono da destra; "Tipo Impianto" è il primo campo che vale
"Stradale" o "Autostradale", "Bandiera" quello che lo precede e "Gestore" tutto ciò che sta prima.
"""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
DIR_ANAGRAFICA = ["Anagrafica primo trimestre", "Anagrafica secondo trimestre"]
DIR_PREZZI = ["Prezzi primo trimestre", "Prezzi secondo trimestre"]
DIR_INTERMEDI = BASE / "dati_intermedi"
DIR_OUTPUT = BASE / "output"
FILE_XLSX = DIR_OUTPUT / "prezzi_medi_carburanti_S1_2026.xlsx"

TIPI_IMPIANTO = {"Stradale", "Autostradale"}

# Regole di calcolo dei prezzi medi
PREZZO_MIN, PREZZO_MAX = 1.0, 3.0
BANDIERE_DISTINTE = ["Agip Eni", "Api-Ip", "Pompe Bianche", "Q8", "Esso", "Tamoil"]
GRANDE_DISTRIBUZIONE = ["CONAD", "Enercoop", "Auchan", "Iper Station", "Carrefour", "COOP"]
ORDINE_OPERATORI = BANDIERE_DISTINTE + ["Grande distribuzione", "Altri operatori"]
ORDINE_PRODOTTI = ["Benzina", "Diesel"]
PRICE_CAP = {"Benzina": 1.99, "Diesel": 2.19}  # €/l, annunciato da Eni, IP e Q8
COLONNE_ANAGRAFICA = ["idImpianto", "Gestore", "Bandiera", "Tipo Impianto",
                      "Comune", "Provincia", "Latitudine", "Longitudine"]


def data_da_nome(path):
    """anagrafica_impianti_attivi-20260101.csv -> date(2026, 1, 1)"""
    return datetime.strptime(Path(path).stem.split("-")[-1], "%Y%m%d").date()


def file_anagrafica():
    return sorted((f for d in DIR_ANAGRAFICA for f in (BASE / d).glob("*.csv")), key=data_da_nome)


def file_prezzi():
    return sorted((f for d in DIR_PREZZI for f in (BASE / d).glob("*.csv")), key=data_da_nome)


def leggi_anagrafica(path):
    righe = []
    with open(path, encoding="utf-8", errors="replace") as fh:
        fh.readline()  # "Estrazione del ..."
        fh.readline()  # intestazione
        for riga in fh:
            campi = riga.rstrip("\r\n").split("|")
            if len(campi) < 10:
                continue
            tipo = next((i for i, c in enumerate(campi[3:-6], start=3)
                         if c.strip() in TIPI_IMPIANTO), 3)
            gestore = "|".join(campi[1:tipo - 1]).strip()
            righe.append([campi[0], gestore, campi[tipo - 1], campi[tipo]] + campi[-4:])
    df = pd.DataFrame(righe, columns=COLONNE_ANAGRAFICA)
    df["Bandiera"] = df["Bandiera"].str.strip()
    df["Tipo Impianto"] = df["Tipo Impianto"].str.strip()
    return df


def leggi_prezzi(path):
    return pd.read_csv(path, sep="|", skiprows=1, dtype={"idImpianto": str, "descCarburante": str,
                                                         "isSelf": "Int8", "dtComu": str},
                       encoding="utf-8", encoding_errors="replace")


def operatore(bandiera):
    if bandiera in BANDIERE_DISTINTE:
        return bandiera
    if bandiera in GRANDE_DISTRIBUZIONE:
        return "Grande distribuzione"
    return "Altri operatori"


def carica_dati():
    """Dataset di 02_prepara_dati.py con le regole di calcolo applicate:
    solo self service, scartati i prezzi fuori da PREZZO_MIN-PREZZO_MAX (salvati in
    output/prezzi_scartati.csv), colonna "operatore" con il raggruppamento delle bandiere."""
    d = pd.read_pickle(DIR_INTERMEDI / "prezzi_benzina_gasolio.pkl")
    d = d[d["isSelf"] == 1]
    anomali = d[(d["prezzo"] < PREZZO_MIN) | (d["prezzo"] > PREZZO_MAX)]
    DIR_OUTPUT.mkdir(exist_ok=True)
    anomali.to_csv(DIR_OUTPUT / "prezzi_scartati.csv", index=False)
    print(f"Prezzi self: {len(d):,}  |  scartati fuori da {PREZZO_MIN}-{PREZZO_MAX} €/l: {len(anomali):,}")
    d = d.drop(anomali.index)

    mappa = {b: operatore(b) for b in d["bandiera"].cat.categories}
    d["operatore"] = d["bandiera"].astype(str).map(mappa)
    return d


# Stile Excel (preferenze XLSX di Riccardo): Arial 10, intestazione bold e congelata,
# testo a sinistra e numeri a destra. I prezzi dei carburanti si leggono al millesimo: 3 decimali.
FONT_DATI = Font(name="Arial", size=10)
FONT_INTESTAZIONE = Font(name="Arial", size=10, bold=True)
SINISTRA, DESTRA = Alignment(horizontal="left"), Alignment(horizontal="right")
FORMATI = {"data": "DD/MM/YYYY", "prezzo": "0.000", "intero": "0", "percentuale": "0.0", "codice": "@"}
TIPI_TESTO = {"testo", "codice"}


def scrivi_foglio(ws, df, tipi):
    """tipi: per ogni colonna 'data', 'testo', 'codice' (testo, es. codici Istat),
    'prezzo', 'intero' o 'percentuale'."""
    for c, (nome, tipo) in enumerate(zip(df.columns, tipi), start=1):
        cella = ws.cell(row=1, column=c, value=nome)
        cella.font = FONT_INTESTAZIONE
        cella.alignment = SINISTRA if tipo in TIPI_TESTO else DESTRA
    for r, riga in enumerate(df.itertuples(index=False), start=2):
        for c, (valore, tipo) in enumerate(zip(riga, tipi), start=1):
            if tipo == "data":
                valore = valore.to_pydatetime().date()
            elif tipo in ("prezzo", "percentuale"):
                valore = float(valore)
            elif tipo == "intero":
                valore = int(valore)
            elif tipo in TIPI_TESTO:
                valore = str(valore)
            cella = ws.cell(row=r, column=c, value=valore)
            cella.font = FONT_DATI
            cella.alignment = SINISTRA if tipo in TIPI_TESTO else DESTRA
            if tipo in FORMATI:
                cella.number_format = FORMATI[tipo]
    for c, nome in enumerate(df.columns, start=1):
        larghezza = max([len(str(nome))] + [len(str(v)) for v in df.iloc[:, c - 1]])
        ws.column_dimensions[get_column_letter(c)].width = max(12, larghezza + 3)
    ws.freeze_panes = "A2"
