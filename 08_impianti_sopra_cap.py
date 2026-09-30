"""
Superamenti del price cap (1,99 €/l benzina, 2,19 €/l diesel) calcolati sui singoli distributori,
prezzi self, impianti stradali, primo semestre 2026. Un prezzo uguale al cap non conta come superamento.

Oltre al filtro generale 1-3 €/l, qui si escludono i prezzi pari o superiori a 2,90 €/l: nel semestre
sono valori segnaposto (soprattutto 2,999 €/l, spesso uguale per benzina e diesel lo stesso giorno),
mentre i prezzi reali più alti, ad esempio nelle isole minori, restano sotto 2,80 €/l.
Una soglia relativa (es. +30% sulla mediana del giorno) non è usata perché scarterebbe proprio quei
prezzi reali (Pantelleria, Lampedusa, Capraia, Capri).

Aggiunge (o sostituisce) due fogli nel file Excel:
- "Impianti sopra il cap": per operatore e prodotto, impianti monitorati, impianti sopra il cap almeno
  un giorno e quota, giornate-impianto sopra il cap e quota, giorni medi sopra il cap per impianto
  (fra quelli che lo hanno superato almeno una volta);
- "Impianti sopra cap per giorno" (formato verticale): per giorno, operatore e prodotto,
  impianti con prezzo, impianti sopra il cap e quota.
I prezzi esclusi sono salvati in output/prezzi_segnaposto_esclusi.csv.
"""
import pandas as pd
from openpyxl import load_workbook

from comune import DIR_OUTPUT, FILE_XLSX, ORDINE_OPERATORI, ORDINE_PRODOTTI, PRICE_CAP, carica_dati, scrivi_foglio

SOGLIA_SEGNAPOSTO = 2.90
FOGLIO_SINTESI = "Impianti sopra il cap"
FOGLIO_GIORNI = "Impianti sopra cap per giorno"


def ordina(df, colonne):
    df = df.copy()
    df["Operatore"] = pd.Categorical(df["Operatore"], ORDINE_OPERATORI, ordered=True)
    df["Prodotto"] = pd.Categorical(df["Prodotto"], ORDINE_PRODOTTI, ordered=True)
    df = df.sort_values(colonne)
    df["Operatore"], df["Prodotto"] = df["Operatore"].astype(str), df["Prodotto"].astype(str)
    return df


def main():
    d = carica_dati()
    segnaposto = d[d["prezzo"] >= SOGLIA_SEGNAPOSTO]
    segnaposto.to_csv(DIR_OUTPUT / "prezzi_segnaposto_esclusi.csv", index=False)
    print(f"Prezzi segnaposto esclusi (>= {SOGLIA_SEGNAPOSTO} €/l): {len(segnaposto):,} "
          f"da {segnaposto['idImpianto'].nunique()} impianti")
    d = d.drop(segnaposto.index)

    d["Prodotto"] = d["prodotto"].astype(str)
    d["Operatore"] = d["operatore"]
    d["sopra"] = d["prezzo"] > d["Prodotto"].map(PRICE_CAP) + 1e-9
    chiave = ["Operatore", "Prodotto"]

    # Sintesi sul semestre
    per_impianto = d.groupby(chiave + ["idImpianto"])["sopra"].sum().rename("giorni_sopra").reset_index()
    sintesi = per_impianto.groupby(chiave).agg(
        impianti=("idImpianto", "size"),
        impianti_sopra=("giorni_sopra", lambda s: int((s > 0).sum())),
        giorni_medi=("giorni_sopra", lambda s: s[s > 0].mean() if (s > 0).any() else 0.0),
    )
    giornate = d.groupby(chiave)["sopra"].agg(["size", "sum"])
    sintesi["quota_impianti"] = (sintesi["impianti_sopra"] / sintesi["impianti"] * 100).round(1)
    sintesi["giornate_sopra"] = giornate["sum"].astype(int)
    sintesi["quota_giornate"] = (giornate["sum"] / giornate["size"] * 100).round(1)
    sintesi["giorni_medi"] = sintesi["giorni_medi"].round(1)
    sintesi = ordina(sintesi.reset_index(), chiave)[
        ["Operatore", "Prodotto", "impianti", "impianti_sopra", "quota_impianti",
         "giornate_sopra", "quota_giornate", "giorni_medi"]]
    sintesi.columns = ["Operatore", "Prodotto", "Impianti monitorati", "Impianti sopra il cap almeno un giorno",
                       "Quota impianti sopra il cap (%)", "Giornate-impianto sopra il cap",
                       "Quota giornate-impianto sopra il cap (%)", "Giorni medi sopra il cap per impianto"]

    # Dettaglio giornaliero
    giorni = d.groupby(["data"] + chiave)["sopra"].agg(["size", "sum"]).reset_index()
    giorni["quota"] = (giorni["sum"] / giorni["size"] * 100).round(1)
    giorni = ordina(giorni, ["data"] + chiave)
    giorni.columns = ["Data", "Operatore", "Prodotto", "Impianti con prezzo", "Impianti sopra il cap",
                      "Quota impianti sopra il cap (%)"]
    assert len(giorni) == 181 * len(ORDINE_OPERATORI) * len(ORDINE_PRODOTTI), "Combinazioni mancanti"
    assert (giorni.groupby(["Operatore", "Prodotto"])["Impianti sopra il cap"].sum().values
            == sintesi.set_index(["Operatore", "Prodotto"]).loc[
                list(giorni.groupby(["Operatore", "Prodotto"]).groups)]["Giornate-impianto sopra il cap"].values).all()

    wb = load_workbook(FILE_XLSX)
    for nome in (FOGLIO_SINTESI, FOGLIO_GIORNI):
        if nome in wb.sheetnames:
            del wb[nome]
    scrivi_foglio(wb.create_sheet(FOGLIO_SINTESI), sintesi,
                  ["testo", "testo", "intero", "intero", "percentuale", "intero", "percentuale", "percentuale"])
    scrivi_foglio(wb.create_sheet(FOGLIO_GIORNI), giorni,
                  ["data", "testo", "testo", "intero", "intero", "percentuale"])
    wb.save(FILE_XLSX)

    with pd.option_context("display.width", 250, "display.max_columns", None):
        print(sintesi.to_string(index=False))
        picco = giorni.loc[giorni.groupby(["Operatore", "Prodotto"])["Quota impianti sopra il cap (%)"].idxmax()]
        print("\nGiorno con la quota più alta di impianti sopra il cap:\n", picco.to_string(index=False))
    print(f"\nAggiunti i fogli '{FOGLIO_SINTESI}' e '{FOGLIO_GIORNI}' a {FILE_XLSX}")


if __name__ == "__main__":
    main()
