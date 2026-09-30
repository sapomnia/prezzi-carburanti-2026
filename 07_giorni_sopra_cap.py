"""
Per ciascun operatore, numero di giorni del primo semestre 2026 in cui il prezzo medio self
(foglio "Per operatore", arrotondato al millesimo) è stato superiore al price cap:
1,99 €/l per la benzina e 2,19 €/l per il diesel. Un prezzo uguale al cap non conta come superamento.

Aggiunge (o sostituisce) il foglio "Giorni sopra il price cap" nel file Excel.
"""
import pandas as pd
from openpyxl import load_workbook

from comune import FILE_XLSX, ORDINE_OPERATORI, ORDINE_PRODOTTI, PRICE_CAP, scrivi_foglio

NOME_FOGLIO = "Giorni sopra il price cap"


def etichetta(prodotto):
    return f"Giorni {prodotto.lower()} sopra {PRICE_CAP[prodotto]:.2f} €/l".replace(".", ",")


def main():
    per_op = pd.read_excel(FILE_XLSX, sheet_name="Per operatore")
    per_op["cap"] = per_op["Prodotto"].map(PRICE_CAP)
    per_op["sopra"] = per_op["Prezzo medio (€/l)"] > per_op["cap"] + 1e-9

    giorni = (per_op.groupby(["Operatore", "Prodotto"])["sopra"].sum().unstack("Prodotto")
              .reindex(index=ORDINE_OPERATORI, columns=ORDINE_PRODOTTI).fillna(0).astype(int))
    assert per_op.groupby(["Operatore", "Prodotto"]).size().eq(181).all(), "Giorni mancanti"
    out = giorni.reset_index()
    out.columns = ["Operatore"] + [etichetta(p) for p in ORDINE_PRODOTTI]

    wb = load_workbook(FILE_XLSX)
    if NOME_FOGLIO in wb.sheetnames:
        del wb[NOME_FOGLIO]
    scrivi_foglio(wb.create_sheet(NOME_FOGLIO), out, ["testo", "intero", "intero"])
    wb.save(FILE_XLSX)

    print(out.to_string(index=False))
    sopra = per_op[per_op["sopra"]]
    print("\nPrimo e ultimo giorno sopra il cap:")
    print(sopra.groupby(["Prodotto", "Operatore"])["Data"].agg(["min", "max"]).to_string())
    print(f"\nAggiunto il foglio '{NOME_FOGLIO}' a {FILE_XLSX}")


if __name__ == "__main__":
    main()
